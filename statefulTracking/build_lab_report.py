from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "output" / "Stateful_Web_Tracking_Lab_Report.docx"
SCREENSHOTS = ROOT / "Screenshots"
REPORT_ASSETS = ROOT / "tmp" / "report_assets"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_code_block(doc, code):
    lines = code.strip("\n").splitlines()
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.55)
    set_cell_shading(cell, "F3F3F3")
    set_cell_border(cell, color="D9D9D9", size="4")
    set_cell_margins(cell, top=100, start=140, bottom=100, end=140)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Courier New")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Courier New")
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(25, 25, 25)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(doc, image_path, caption, width=6.45):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Inches(width))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", caption)

    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(8)
    cap.add_run(caption)


def prepare_report_assets():
    """Create a focused crop of the cookie evidence for legibility in the report."""
    REPORT_ASSETS.mkdir(parents=True, exist_ok=True)
    source = SCREENSHOTS / "Screenshot 2026-10-02 at 18.13.40.png"
    target = REPORT_ASSETS / "challenge2_publisher_cookie.png"
    with Image.open(source) as image:
        image.crop((0, 745, image.width, image.height)).save(target)
    return target


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(20, 20, 20)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(6)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 16, 14, 6),
        ("Heading 2", 12.5, 10, 4),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Aptos"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(55, 55, 55)

    footer = section.footer
    footer.is_linked_to_previous = False
    table = footer.add_table(rows=1, cols=2, width=Inches(6.9))
    table.autofit = False
    table.columns[0].width = Inches(5.8)
    table.columns[1].width = Inches(1.1)
    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    run = left.add_run("Stateful Web Tracking Lab")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(90, 90, 90)
    right = table.cell(0, 1).paragraphs[0]
    right.paragraph_format.space_after = Pt(0)
    add_page_number(right)


def add_status_table(doc):
    rows = [
        ("1", "Independent publishers", "Met", "Flask applications run on ports 8001 and 8002."),
        ("2", "Tracker embedded", "Met", "Both publisher pages load tracker-one.test:9000 in an iframe."),
        ("3", "Unique identifier", "Partial", "Tracker 1 generates aid, but the distinct-host test does not retain it across iframe reloads."),
        ("4", "Identifier cookie", "Partial", "The server sends Set-Cookie, but Chrome blocks it in the current cross-site HTTP context."),
        ("5", "Request recording", "Met", "Every request and its headers are printed in the tracker terminal during the session."),
        ("6", "Identifier linked to publisher and page", "Partial", "Query parameters identify the publisher and page, but structured profile output is deferred."),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.38, 1.45, 0.68, 4.05]
    headers = ["No", "Requirement", "Status", "Current evidence"]
    for index, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = Inches(width)
        set_cell_shading(cell, "404040")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)
    set_repeat_table_header(table.rows[0])

    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, text, width) in enumerate(zip(cells, row_data, widths)):
            cell.width = Inches(width)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, "F7F7F7")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.size = Pt(8.6)
            if index == 2:
                run.bold = True
    return table


def build_report():
    doc = Document()
    style_document(doc)
    challenge2_cookie = prepare_report_assets()

    title = doc.add_paragraph(style="Title")
    title.add_run("Stateful Web Tracking Lab Report")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(4)
    run = subtitle.add_run("Warm Up, Challenge 1 and Challenge 2")
    run.bold = True
    run.font.size = Pt(13)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    meta_run = meta.add_run("Web Privacy TP1  |  2 October 2026")
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = RGBColor(80, 80, 80)

    intro = doc.add_paragraph()
    intro.add_run("Summary. ").bold = True
    intro.add_run(
        "The warm-up confirmed how Flask exposes HTTP headers and how a server assigns an aid cookie. "
        "Challenge 1 then used two independent publishers that embed the same tracker. The tracker logic works "
        "when the browser accepts the cookie, but changing from 127.0.0.1 to separate .test hosts made the iframe "
        "cross-site. Chrome therefore rejected the cookie because it has no SameSite=None attribute and the lab is using HTTP. "
        "Challenge 2 instead uses a publisher-owned first-party analytics cookie and explicitly sends the identifier to an analytics endpoint."
    )

    doc.add_heading("Warm Up", level=1)
    doc.add_heading("Inspecting HTTP Headers", level=2)
    doc.add_paragraph(
        "The Flask application prints the browser request headers and the response headers. The Network panel and "
        "the terminal expose the same exchange from client-side and server-side viewpoints. The request includes the "
        "Host and User-Agent headers, while the response includes content metadata and any custom or cookie headers added by Flask."
    )

    doc.add_heading("Creating and Reusing the Cookie", level=2)
    add_code_block(
        doc,
        '''aid = request.cookies.get("aid")
is_new = aid is None

if is_new:
    aid = secrets.token_hex(8)

response = make_response(render_template("index.html"))
if is_new:
    response.set_cookie(key="aid", value=aid)''',
    )
    doc.add_paragraph(
        "On the first request, no aid cookie exists, so the server generates a 16-character hexadecimal identifier and "
        "returns it in Set-Cookie. Later requests send the same value in the Cookie request header, allowing the server "
        "to recognize the browser."
    )

    answer = doc.add_paragraph()
    answer.add_run("Browser restart question. ").bold = True
    answer.add_run(
        "The default Flask cookie has no Max-Age or Expires attribute, so it is a session cookie. It normally remains "
        "available during reloads in the same browser session, but it is not designed to survive a complete browser-session end. "
        "A browser that restores the previous session may restore it as an implementation detail."
    )

    doc.add_page_break()
    doc.add_heading("Challenge 1 Third Party Tracking", level=1)
    doc.add_paragraph(
        "The TP requires two publishers, one shared tracker, a browser identifier stored in a cookie, request logging, "
        "and an association between the identifier and the publisher page. The chronological profile presentation is "
        "reserved for a later report update."
    )
    add_status_table(doc)

    doc.add_heading("Publisher and Tracker Implementation", level=2)
    doc.add_paragraph(
        "Publisher 1 runs on port 8001 and Publisher 2 runs on port 8002. Each page supplies its publisher name and "
        "page name in the tracker iframe URL."
    )
    add_code_block(
        doc,
        '''<!-- Publisher 1 -->
<iframe src="http://tracker-one.test:9000/
?publisher=publisher-one&page=Travel"></iframe>

<!-- Publisher 2 -->
<iframe src="http://tracker-one.test:9000/
?publisher=publisher-two&page=Technology"></iframe>''',
    )
    add_code_block(
        doc,
        '''publisher = request.args.get("publisher", "unknown publisher")
page = request.args.get("page", "unknown page")
aid = request.cookies.get("aid")

if aid is None:
    aid = secrets.token_hex(8)
    response.set_cookie(key="aid", value=aid)

print(request.headers)
print(response.headers)''',
    )
    doc.add_paragraph(
        "The query parameters identify the source page, and the terminal output records each received request. "
        "The current implementation does not yet store structured history or automatically print a chronological profile."
    )

    doc.add_page_break()
    doc.add_heading("Observed Cookie Behaviour", level=1)
    doc.add_paragraph(
        "The following earlier test used 127.0.0.1 for the publishers and tracker. Ports create different origins, but "
        "they do not create different sites for SameSite cookie evaluation. The cookie was therefore accepted and reused."
    )
    add_figure(
        doc,
        SCREENSHOTS / "Screenshot 2026-10-01 at 10.22.06.png",
        "Figure 1  The first tracker request has no aid cookie, so Tracker 1 returns Set-Cookie with aid=860cceb220522113.",
    )
    add_figure(
        doc,
        SCREENSHOTS / "Screenshot 2026-10-01 at 10.22.24.png",
        "Figure 2  A later request from Publisher 2 sends the same aid value, demonstrating identifier reuse.",
    )

    doc.add_page_break()
    doc.add_heading("Effect of the Hostname Change", level=1)
    doc.add_paragraph(
        "The current publisher URLs use publisher-one.test and publisher-two.test, while the iframe uses tracker-one.test. "
        "These hostnames are different sites. The live tracker response sends only the following cookie attributes:"
    )
    add_code_block(doc, "Set-Cookie: aid=<random identifier>; Path=/")
    doc.add_paragraph(
        "Because SameSite is omitted, Chrome treats the cookie as SameSite=Lax. A Lax cookie cannot be set from this "
        "cross-site iframe. Browser diagnostics reported isSameSite: false, no associated cookies, and the rejection "
        "reason SameSiteUnspecifiedTreatedAsLax. The next iframe request therefore contains no aid, and the tracker creates a new value."
    )

    comparison = doc.add_table(rows=1, cols=3)
    comparison.alignment = WD_TABLE_ALIGNMENT.CENTER
    comparison.autofit = False
    compare_headers = ("Configuration", "Browser classification", "Result")
    compare_rows = (
        ("127.0.0.1 on different ports", "Same-site", "aid is accepted and reused"),
        ("publisher-one.test and tracker-one.test", "Cross-site", "Current aid is rejected in the iframe"),
    )
    widths = (2.35, 1.55, 2.65)
    for index, (cell, text, width) in enumerate(zip(comparison.rows[0].cells, compare_headers, widths)):
        cell.width = Inches(width)
        set_cell_shading(cell, "404040")
        set_cell_border(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)
    for row_index, row_data in enumerate(compare_rows):
        cells = comparison.add_row().cells
        for cell, text, width in zip(cells, row_data, widths):
            cell.width = Inches(width)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, "F7F7F7")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(9)

    doc.add_heading("Challenge 1 Result", level=1)
    doc.add_paragraph(
        "The warm-up demonstrated stateful recognition with a session cookie. Challenge 1 established the required "
        "publisher and tracker applications and showed that one tracker identifier can follow visits across publishers "
        "when the cookie is accepted. The distinct-host test exposes a modern browser restriction: SameSite=None alone "
        "is insufficient because it also requires Secure, and Secure requires HTTPS for these custom hostnames. The code "
        "is intentionally left unchanged for this part of the lab."
    )

    doc.add_page_break()
    doc.add_heading("Challenge 2 First Party Analytics", level=1)
    doc.add_paragraph(
        "Challenge 2 changes where the identifier is stored and how it reaches the analytics service. The external "
        "analytics.js file executes inside each publisher page, creates an analytics_id cookie for that publisher, and "
        "uses fetch() to send the identifier and page data to analytics.test. Because the cookie belongs to the publisher, "
        "the two publisher hostnames receive separate identifiers."
    )

    question = doc.add_paragraph()
    question.add_run("Does the analytics server automatically receive the publisher cookie? ").bold = True
    question.add_run(
        "No. A cookie scoped to publisher-one.test is sent automatically only to publisher-one.test. JavaScript must read "
        "the value and include it in the cross-origin request body. The browser still applies CORS to that request."
    )

    add_code_block(
        doc,
        '''const analyticsId = readCookie("analytics_id") || generateIdentifier();
document.cookie = `analytics_id=${analyticsId}; Max-Age=31536000; Path=/; SameSite=Lax`;

fetch("http://analytics.test:9100/collect", {
  method: "POST",
  headers: { "Content-Type": "application/json" },
  body: JSON.stringify({ analytics_id: analyticsId, publisher, page, page_title, timestamp })
});''',
    )

    add_figure(
        doc,
        challenge2_cookie,
        "Figure 3  Publisher 1 stores analytics_id as a first-party cookie scoped to publisher-one.test.",
        width=6.45,
    )
    add_figure(
        doc,
        SCREENSHOTS / "Screenshot 2026-10-02 at 18.17.34.png",
        "Figure 4  The collect request body explicitly carries the identifier, publisher, page, title and timestamp.",
        width=6.45,
    )

    doc.add_page_break()
    doc.add_heading("Challenge 2 Request Flow and Evidence", level=1)
    preflight = doc.add_paragraph()
    preflight.add_run("CORS preflight. ").bold = True
    preflight.add_run(
        "The JSON POST crosses from publisher-one.test to analytics.test, so the browser first sends an OPTIONS request. "
        "A 204 response authorizes the origin, POST method and Content-Type header; the browser then sends the actual POST, "
        "which also returns 204. The OPTIONS request is a permission check, not a second analytics event."
    )

    add_figure(
        doc,
        SCREENSHOTS / "Screenshot 2026-10-02 at 18.14.13.png",
        "Figure 5  Analytics server log: script delivery followed by successful OPTIONS and POST requests to /collect.",
        width=5.75,
    )
    add_figure(
        doc,
        SCREENSHOTS / "Screenshot 2026-10-02 at 18.18.13.png",
        "Figure 6  Publisher 1 receives analytics_id in its Cookie request header, confirming first-party persistence.",
        width=5.35,
    )

    doc.add_heading("Challenge 2 Status", level=2)
    doc.add_paragraph(
        "The implemented scope is complete for identifier creation, first-party cookie persistence, event construction, "
        "CORS handling and POST collection. Storage of collected events and chronological profile reconstruction remain "
        "intentionally deferred, corresponding to the omitted requirements 6 and 7. Unlike Challenge 1, this approach "
        "keeps a separate identifier in each publisher's cookie jar; linking those identifiers would require a later mechanism."
    )

    source = doc.add_paragraph()
    source.paragraph_format.space_before = Pt(10)
    source.paragraph_format.space_after = Pt(0)
    run = source.add_run("Lab source: TP Privacy, TP1 Stateful Tracking, sections 3 through 5.")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(90, 90, 90)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Stateful Web Tracking Lab Report"
    doc.core_properties.subject = "Warm Up, Challenge 1 Third Party Tracking and Challenge 2 Analytics"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "web privacy, cookies, SameSite, third-party tracking, analytics, CORS"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
